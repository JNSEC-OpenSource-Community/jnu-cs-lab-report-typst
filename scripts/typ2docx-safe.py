#!/usr/bin/env python3
"""Use typ2docx's PDF-to-DOCX engine and restore the editable header layout."""

import argparse
import copy
import re
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from pdf2docx import parse
from pygments import lex
from pygments.lexers import TextLexer, get_lexer_by_name, guess_lexer
from pygments.token import Comment, Keyword, Literal, String


REPORT_COLUMNS_CM = (
    2.095,
    0.635, 0.635, 0.635, 0.635,
    0.635, 0.635, 0.635, 0.635,
    1.270, 0.635, 1.270, 1.905, 1.280, 2.847,
)
HEADER_ROW_HEIGHT_CM = 0.64
HEADER_REQUIREMENTS_GAP_CM = 0.10
HEADER_TEXT_SIZE = 10.5
DEFAULT_HEADER_LABEL_GAP_CM = 0.75
HEADER_FIELDS = (
    "course_name",
    "experiment_name",
    "experiment_date",
    "class_name",
    "student_name",
    "student_id",
    "instrument_id",
    "report_requirements",
)
BODY_HEADINGS = ("实验目的：", "实验内容：", "程序清单：", "运行情况：", "实验体会：")
BODY_TEXT_SIZE = 10.5
BODY_HEADING_SIZE = 14
BODY_MIN_HEIGHT_CM = 16.174
LANGUAGE_ALIASES = {
    "py": "python",
    "js": "javascript",
    "jsx": "javascript",
    "ts": "typescript",
    "tsx": "typescript",
    "c++": "cpp",
    "h++": "cpp",
    "cc": "cpp",
    "sh": "bash",
    "shell": "bash",
    "zsh": "bash",
    "yml": "yaml",
    "md": "markdown",
    "text": "text",
}


def _read_typst_fields(source: Path):
    text = source.read_text(encoding="utf-8")
    values = {}
    for name in HEADER_FIELDS:
        match = re.search(
            rf"(?m)^\s*#let\s+{re.escape(name)}\s*=\s*\[([^\]]*)\]", text
        )
        values[name] = match.group(1).strip() if match else ""
    gap_match = re.search(
        r"班\s*#h\(\s*([0-9]+(?:\.[0-9]+)?)\s*cm\s*\)\s*级",
        text,
    )
    values["header_label_gap_cm"] = (
        float(gap_match.group(1)) if gap_match else DEFAULT_HEADER_LABEL_GAP_CM
    )
    return values


def _read_program_blocks(source: Path):
    text = source.read_text(encoding="utf-8")
    match = re.search(
        r"(?ms)^\s*#let\s+program_list\s*=\s*\[(.*?)^\s*\]\s*(?=^\s*#let\s+run_status)",
        text,
    )
    if not match:
        return []

    body = match.group(1)
    positioned_blocks = []

    for fenced in re.finditer(
        r"(?ms)^[^\S\n]*```([^\n]*)\n(.*?)^[^\S\n]*```[^\S\n]*$", body
    ):
        language, block = fenced.groups()
        positioned_blocks.append(
            (
                fenced.start(),
                language.strip() or "text",
                textwrap.dedent(block).strip("\n"),
            )
        )

    # Reports may keep long listings in separate source files and include them
    # with Typst's `raw(read(...))`.  Reading those files directly is essential:
    # PDF-to-DOCX text recognition does not reliably preserve code whitespace.
    for included in re.finditer(
        r'''(?ms)#raw\(\s*read\(\s*"([^"]+)"\s*\)\s*,\s*lang:\s*"([^"]+)"\s*,\s*block:\s*true\s*\)''',
        body,
    ):
        relative_path, language = included.groups()
        code_path = source.parent / relative_path
        positioned_blocks.append(
            (
                included.start(),
                language.strip() or "text",
                code_path.read_text(encoding="utf-8"),
            )
        )

    positioned_blocks.sort(key=lambda item: item[0])
    return [(language, block) for _, language, block in positioned_blocks]


def _compact_code_text(text):
    return re.sub(r"\s+", "", text)


def _split_after_nonspace(text, count):
    """Split text immediately after `count` non-whitespace characters."""
    if count <= 0:
        return "", text
    seen = 0
    for index, character in enumerate(text):
        if not character.isspace():
            seen += 1
            if seen == count:
                return text[: index + 1], text[index + 1 :]
    return text, ""


def _document_table_paragraphs(document):
    """Return table paragraphs in visual order without merged-cell duplicates."""
    paragraphs = []
    seen_cells = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen_cells:
                    continue
                seen_cells.add(marker)
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def _append_cell_children(destination, source):
    """Append a cell's paragraphs/drawings while preserving their XML."""
    for child in source._tc:
        if child.tag == qn("w:tcPr"):
            continue
        destination._tc.append(copy.deepcopy(child))


def _collapse_pdf_page_tables(document):
    """Collapse pdf2docx's one-table-per-page output into one report table.

    The Typst report is one 15-column outer table whose header and body rows
    span all columns.  pdf2docx represents each PDF page as a separate table;
    moving those body fragments into the first table restores the intended
    editable structure and keeps the evaluation row attached to the report.
    """
    tables = list(document.tables)
    if len(tables) < 2 or len(tables[0].columns) != 1:
        return
    evaluation_table = tables[-1]
    if len(evaluation_table.columns) != 15 or len(evaluation_table.rows) < 2:
        return

    first = tables[0]
    first_text = " ".join(cell.text for cell in first.rows[0].cells)
    has_header = "课程名称" in first_text
    body_row_index = 1 if has_header else 0
    body_cell = first.rows[body_row_index].cells[0]

    if not has_header and len(first.rows) > 1 and _row_is_empty_continuation(first.rows[1]):
        continuation = first.rows[1]._tr
        continuation.getparent().remove(continuation)

    # Append all continuation-page body content in visual order.  The final
    # table's first row contains the reflection section; its second row is the
    # 15-cell teacher-evaluation row and is handled separately below.
    for table in tables[1:-1]:
        xml_cells = table.rows[0]._tr.findall(qn("w:tc"))
        if xml_cells:
            source_cell = table.rows[0].cells[0]
            _append_cell_children(body_cell, source_cell)
    reflection_cell = evaluation_table.rows[0]._tr.findall(qn("w:tc"))[0]
    for child in reflection_cell:
        if child.tag != qn("w:tcPr"):
            body_cell._tc.append(copy.deepcopy(child))

    # Expand the first table to the report grid before merging its header and
    # body rows.  The evaluation row is appended after expansion so it keeps
    # its original fifteen cells.
    _remove_vertical_merge(body_cell)
    if has_header:
        _remove_vertical_merge(first.rows[0].cells[0])
    for _ in range(14):
        first.add_column(Cm(0.635))
    if has_header:
        _merge_row(first, 0)
        _merge_row(first, 1)
    else:
        _merge_row(first, 0)
    first._tbl.append(copy.deepcopy(evaluation_table.rows[1]._tr))

    parent = first._tbl.getparent()
    for table in tables[1:]:
        parent.remove(table._tbl)


def _format_code_paragraph(paragraph, first=False, last=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4.25 if first else 0)
    paragraph.paragraph_format.space_after = Pt(4.25 if last else 0)
    paragraph.paragraph_format.line_spacing = Pt(11)
    # Keep each source fragment (normally one or two complete functions)
    # together.  Otherwise Word may leave a lone closing brace in a repeated
    # table-row border at the top of the next page.
    paragraph.paragraph_format.keep_together = True


def _add_highlighted_code(paragraph, code, language):
    for token, value in lex(code, _code_lexer(language, code)):
        run = paragraph.add_run(value)
        _set_code_run_font(run)
        run.font.color.rgb = RGBColor.from_string(_code_color(token))


def _remove_paragraph(paragraph):
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def _element(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_borders(parent, edges):
    border_tag = "w:tcBorders" if parent.tag == qn("w:tcPr") else "w:tblBorders"
    borders = _element(parent, border_tag)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV", "start", "end"):
        settings = edges.get(edge, {"val": "nil"})
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in settings.items():
            node.set(qn(f"w:{key}"), str(value))


def _set_cell_margins(cell, start=0, end=0, top=0, bottom=0):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in list(tc_pr.findall(qn("w:tcMar"))):
        tc_pr.remove(old)
    margins = OxmlElement("w:tcMar")
    tc_pr.append(margins)
    for side, value in (("start", start), ("end", end), ("top", top), ("bottom", bottom)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(round(value)))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _set_run_font(run, size, font="NSimSun"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), font)


def _set_code_run_font(run, size=10.5):
    run.font.name = "Consolas"
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for font_type in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{font_type}"), "Consolas")
    for text_node in run._element.iter(qn("w:t")):
        text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _code_color(token):
    if token in Comment.Preproc:
        return "E02E4D"
    if token in Comment.PreprocFile or token in String:
        return "008000"
    if token in Comment:
        return "008000"
    if token in Keyword or token in Literal.Number:
        return "E02E4D"
    return "000000"


def _code_lexer(language, code=""):
    language = language.strip().split(None, 1)[0].lower() if language.strip() else "text"
    language = LANGUAGE_ALIASES.get(language, language)
    options = {"stripnl": False, "ensurenl": False}
    try:
        return get_lexer_by_name(language, **options)
    except Exception:
        # Unknown labels can still receive highlighting when Pygments can
        # infer a lexer from the source; otherwise keep the plain-text fallback.
        try:
            return guess_lexer(code, **options) if code.strip() else TextLexer(**options)
        except Exception:
            return TextLexer(**options)


def _set_run_bold(run):
    run.bold = True
    r_pr = run._element.get_or_add_rPr()
    bold = r_pr.find(qn("w:b"))
    if bold is None:
        bold = OxmlElement("w:b")
        r_pr.append(bold)
    bold.set(qn("w:val"), "1")


def _set_cell_text(
    cell,
    text,
    size=10.5,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    left_indent_cm=0,
):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    _clear_paragraph(paragraph)
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Cm(left_indent_cm)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(size)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    _set_run_font(run, size, font="SimSun")


def _set_table_widths(table):
    total_twips = sum(round(width * 567) for width in REPORT_COLUMNS_CM)
    tbl_pr = table._tbl.tblPr
    tbl_w = _element(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_twips))
    layout = _element(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    indent = _element(tbl_pr, "w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")

    for index, width in enumerate(REPORT_COLUMNS_CM):
        twips = round(width * 567)
        table.columns[index].width = Cm(width)
        table._tbl.tblGrid.gridCol_lst[index].set(qn("w:w"), str(twips))
        for cell in table.columns[index].cells:
            cell.width = Cm(width)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def _set_table_borders(table):
    border = {"val": "single", "sz": "4", "color": "000000", "space": "0"}
    edges = {
        edge: border
        for edge in ("top", "left", "bottom", "right", "start", "end", "insideH", "insideV")
    }
    _set_borders(table._tbl.tblPr, edges)
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            _set_borders(cell._tc.get_or_add_tcPr(), {
                edge: border
                for edge in ("top", "left", "bottom", "right", "start", "end")
            })


def _clear_cell_content(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)
    assert tc_pr.getparent() is cell._tc


def _set_header_line(paragraph, runs, space_after=0):
    _clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = Pt(13)
    paragraph.paragraph_format.keep_together = True
    for text, underline in runs:
        run = paragraph.add_run(text)
        run.underline = underline
        _set_run_font(run, HEADER_TEXT_SIZE)


def _set_header_text_paragraph(paragraph):
    _clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Cm(HEADER_ROW_HEIGHT_CM)
    paragraph.paragraph_format.keep_together = True


def _add_tab(paragraph, *, underline=False):
    run = paragraph.add_run()
    run.underline = underline
    _set_run_font(run, HEADER_TEXT_SIZE)
    run._r.append(OxmlElement("w:tab"))


def _add_header_label(paragraph, text, *, spaced=False, gap_cm=DEFAULT_HEADER_LABEL_GAP_CM):
    if spaced:
        first, second = text
        first_run = paragraph.add_run(first)
        _set_run_font(first_run, HEADER_TEXT_SIZE)
        _add_tab(paragraph)
        second_run = paragraph.add_run(second)
        _set_run_font(second_run, HEADER_TEXT_SIZE)
        return

    run = paragraph.add_run(text)
    _set_run_font(run, HEADER_TEXT_SIZE)


def _add_header_field(paragraph, text):
    # A thin leading space mirrors Typst's 0.05 cm field inset.  The final tab
    # is underlined too, so the rule always reaches the fixed field boundary.
    run = paragraph.add_run("\u2009" + text)
    run.underline = True
    _set_run_font(run, HEADER_TEXT_SIZE)
    _add_tab(paragraph, underline=True)


def _set_header_tab_stops(paragraph, stops_cm):
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.clear_all()
    for position in stops_cm:
        tab_stops.add_tab_stop(Cm(position), WD_TAB_ALIGNMENT.LEFT)


def _write_header_line(paragraph, fields, *, gap_cm):
    """Write one fixed-position metadata line without a nested table."""
    _set_header_text_paragraph(paragraph)

    internal_label_stops = []
    glyph_width_cm = HEADER_TEXT_SIZE * 2.54 / 72
    for label_start, _label, spaced, _value, _field_end, _next_start in fields:
        if spaced:
            internal_label_stops.append(label_start + glyph_width_cm + gap_cm)

    used_boundaries = {
        coordinate
        for _label_start, _label, _spaced, _value, field_end, next_start in fields
        for coordinate in (field_end, next_start)
        if coordinate is not None
    }
    field_starts = {label_start + 1.50 for label_start, *_rest in fields}
    _set_header_tab_stops(
        paragraph,
        sorted(internal_label_stops + list(used_boundaries | field_starts)),
    )

    for index, (label_start, label, spaced, value, field_end, next_start) in enumerate(fields):
        if index:
            _add_tab(paragraph)
        _add_header_label(paragraph, label, spaced=spaced, gap_cm=gap_cm)
        _add_tab(paragraph)
        _add_header_field(paragraph, value)
        if next_start is None:
            # LibreOffice collapses a trailing tab at end-of-paragraph.  A
            # zero-width text node keeps the final field underline expanded.
            tail = paragraph.add_run("\u200b")
            _set_run_font(tail, HEADER_TEXT_SIZE)
            break


def _append_header_paragraph(cell):
    paragraph = OxmlElement("w:p")
    cell._tc.append(paragraph)
    return Paragraph(paragraph, cell)


def _write_header_cell(header_cell, values):
    _clear_cell_content(header_cell)
    _set_cell_margins(
        header_cell,
        start=round(0.20 * 567),
        end=round(0.20 * 567),
        top=round(0.20 * 567),
        bottom=0,
    )

    # The two metadata rows are ordinary paragraphs.  Fixed tab stops preserve
    # the Typst column geometry without creating hidden table gridlines in Word
    # or ONLYOFFICE.
    row_1 = _append_header_paragraph(header_cell)
    _write_header_line(
        row_1,
        (
            (0.00, "课程名称", False, values["course_name"], 3.85, 3.97),
            (3.97, "实验名称", False, values["experiment_name"], 9.02, 9.14),
            (9.14, "实验日期", False, values["experiment_date"], 12.64, None),
        ),
        gap_cm=values["header_label_gap_cm"],
    )

    row_2 = _append_header_paragraph(header_cell)
    _write_header_line(
        row_2,
        (
            (0.00, ("班", "级"), True, values["class_name"], 3.85, 3.97),
            (3.97, ("姓", "名"), True, values["student_name"], 9.02, 9.14),
            (9.14, ("学", "号"), True, values["student_id"], 12.64, 12.76),
            # The outer cell has 0.20 cm padding on each side, leaving 15.98 cm.
            # Keep the final tab inside that usable width so blank fields retain
            # their complete underline in LibreOffice and ONLYOFFICE.
            (12.76, "仪器编号", False, values["instrument_id"], 15.98, None),
        ),
        gap_cm=values["header_label_gap_cm"],
    )

    requirements = _append_header_paragraph(header_cell)
    requirements.alignment = WD_ALIGN_PARAGRAPH.LEFT
    requirements.paragraph_format.space_before = Cm(HEADER_REQUIREMENTS_GAP_CM)
    requirements.paragraph_format.space_after = Pt(0)
    requirements.paragraph_format.line_spacing = Pt(9)
    # Typst collapses consecutive source whitespace in text markup; Word does
    # not.  Normalize it here so the requirement line keeps the same spacing
    # and remains on one line as it does in the Typst preview.
    requirement_text = re.sub(r"\s+", " ", values["report_requirements"]).strip()
    run = requirements.add_run(requirement_text)
    _set_run_font(run, 9)


def _set_header_row_height(table):
    table.rows[0].height = Cm(2.1)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _merge_row(table, row_index):
    row = table.rows[row_index]
    xml_cells = row._tr.findall(qn("w:tc"))
    if len(xml_cells) <= 1:
        return row.cells[0]

    merged = row.cells[0]
    for index in range(1, len(table.columns)):
        merged = merged.merge(row.cells[index])
    return merged


def _insert_row_before(table, row_index):
    row = table.add_row()
    row_xml = row._tr
    first_xml = table.rows[row_index]._tr
    row_xml.getparent().remove(row_xml)
    first_xml.addprevious(row_xml)
    return table.rows[row_index]


def _remove_vertical_merge(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    v_merge = tc_pr.find(qn("w:vMerge"))
    if v_merge is not None:
        tc_pr.remove(v_merge)


def _row_is_empty_continuation(row):
    xml_cells = row._tr.findall(qn("w:tc"))
    if len(xml_cells) != 1:
        return False
    tc = xml_cells[0]
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None or tc_pr.find(qn("w:vMerge")) is None:
        return False
    return not any(tc.iter(qn("w:t")))


def _rebuild_missing_header(table, values):
    """Restore the nested Typst header that pdf2docx may omit."""
    if not table.rows:
        return False

    first_text = " ".join(cell.text for cell in table.rows[0].cells)
    if not any(heading in first_text for heading in BODY_HEADINGS):
        return False

    # pdf2docx puts the body in the first row and adds an empty continuation
    # row for the vertically merged cell.  Make it one normal body row before
    # inserting the missing header row.
    if len(table.rows) > 1 and _row_is_empty_continuation(table.rows[1]):
        continuation = table.rows[1]._tr
        continuation.getparent().remove(continuation)

    body_row = table.rows[0]
    _remove_vertical_merge(body_row.cells[0])
    body_row.height = Cm(BODY_MIN_HEIGHT_CM)
    body_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    header_row = _insert_row_before(table, 0)
    header_cell = _merge_row(table, 0)
    _write_header_cell(header_cell, values)
    _set_header_row_height(table)
    return True


def _row_has_multiple_xml_cells(row):
    return len(row._tr.findall(qn("w:tc"))) > 1


def _write_body_cell(body_cell, body_text):
    _set_cell_margins(
        body_cell,
        start=round(0.20 * 567),
        end=round(0.20 * 567),
        top=round(0.34 * 567),
        bottom=round(0.20 * 567),
    )
    paragraph = body_cell.paragraphs[0]
    _clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(14)
    run = paragraph.add_run(body_text)
    _set_run_font(run, BODY_TEXT_SIZE)


def _restore_program_blocks(document, source):
    blocks = _read_program_blocks(source)
    if not blocks or not document.tables:
        return set()

    paragraphs = _document_table_paragraphs(document)
    code_paragraph_ids = set()
    search_start = 0
    for language, block in blocks:
        first_line = next((line.strip() for line in block.splitlines() if line.strip()), "")
        signature = _compact_code_text(first_line)
        block_compact = _compact_code_text(block)
        start_index = next(
            (
                index
                for index in range(search_start, len(paragraphs))
                if signature and signature in _compact_code_text(paragraphs[index].text)
            ),
            -1,
        )
        if start_index < 0 or not block_compact:
            continue

        start_compact = _compact_code_text(paragraphs[start_index].text)
        prefix_count = start_compact.find(signature)
        remaining_count = len(block_compact)
        assignments = []
        valid = True
        for index in range(start_index, len(paragraphs)):
            paragraph = paragraphs[index]
            paragraph_compact = _compact_code_text(paragraph.text)
            current_prefix_count = prefix_count if index == start_index else 0
            available = paragraph_compact[current_prefix_count:]
            code_count = min(len(available), remaining_count)
            consumed_count = len(block_compact) - remaining_count
            if available[:code_count] != block_compact[
                consumed_count : consumed_count + code_count
            ]:
                valid = False
                break
            assignments.append(
                (paragraph, current_prefix_count, code_count, len(available) > code_count)
            )
            remaining_count -= code_count
            if remaining_count == 0:
                break

        if not valid or remaining_count:
            continue

        source_position = 0
        for assignment_index, (paragraph, leading_count, code_count, has_trailing) in enumerate(assignments):
            is_first = assignment_index == 0
            is_last = assignment_index == len(assignments) - 1
            leading, visible = _split_after_nonspace(paragraph.text, leading_count)
            _, trailing = _split_after_nonspace(visible, code_count)

            segment_start = source_position
            seen = 0
            while source_position < len(block) and seen < code_count:
                if not block[source_position].isspace():
                    seen += 1
                source_position += 1
            segment = block[segment_start:source_position]

            if not is_last:
                whitespace_end = source_position
                while whitespace_end < len(block) and block[whitespace_end].isspace():
                    whitespace_end += 1
                separator = block[source_position:whitespace_end]
                if "\n" in separator:
                    source_position += separator.index("\n") + 1
                else:
                    # PDF2DOCX normally splits code on source line boundaries.
                    # If it splits within a line, retain the separating spaces.
                    segment += separator
                    source_position = whitespace_end
            else:
                segment += block[source_position:]
                source_position = len(block)

            _clear_paragraph(paragraph)
            _format_code_paragraph(paragraph, first=is_first, last=is_last)
            if leading:
                run = paragraph.add_run(leading)
                _set_run_font(run, BODY_TEXT_SIZE)
            _add_highlighted_code(paragraph, segment, language)
            if has_trailing and trailing:
                run = paragraph.add_run(trailing)
                size = (
                    BODY_HEADING_SIZE
                    if any(heading in trailing for heading in BODY_HEADINGS)
                    else BODY_TEXT_SIZE
                )
                _set_run_font(run, size)
            code_paragraph_ids.add(id(paragraph._p))

        search_start = start_index + len(assignments)
    return code_paragraph_ids


def _restore_body_font_size(document, code_paragraph_ids):
    for table_index, table in enumerate(document.tables):
        for row in table.rows:
            if any(cell.text.strip() == "教师评价" for cell in row.cells):
                continue
            if any("课程名称" in cell.text for cell in row.cells):
                continue
            for cell in row.cells:
                _set_cell_margins(
                    cell,
                    start=round(0.20 * 567),
                    end=round(0.20 * 567),
                )
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.left_indent = Pt(0)
                    paragraph.paragraph_format.right_indent = Pt(0)
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    # Paragraph proxy objects can be recreated while walking
                    # merged table cells, so their Python ids are not stable.
                    # Keep any paragraph already rebuilt with the code font;
                    # otherwise this pass would turn it back into Times New
                    # Roman and make indentation spaces appear to disappear.
                    if (
                        id(paragraph._p) in code_paragraph_ids
                        or any(run.font.name == "Consolas" for run in paragraph.runs)
                    ):
                        continue
                    for run in paragraph.runs:
                        size = (
                            BODY_HEADING_SIZE
                            if any(heading in run.text for heading in BODY_HEADINGS)
                            else BODY_TEXT_SIZE
                        )
                        _set_run_font(run, size)


def _normalize_code_paragraphs(document):
    """Use one editable, monospaced C style for every listing fragment.

    PDF-to-DOCX may split a raw listing into several table paragraphs and
    assign the surrounding document font to some fragments.  The paragraph
    text itself still contains the source whitespace, so rebuilding each
    listing fragment in place fixes the mixed-font appearance without changing
    pagination or content.
    """
    paragraphs = _document_table_paragraphs(document)
    # The first PDF page can leave the final short static-list function in a
    # split table-row continuation.  Move that whole function to the next
    # listing fragment so Word does not render a lone brace in a border strip.
    for index, paragraph in enumerate(paragraphs[:-1]):
        if "/* 判断空表 */" in paragraph.text and index + 1 < len(paragraphs):
            destination = paragraphs[index + 1]
            moved_text = paragraph.text.strip("\n")
            destination_text = destination.text
            destination_text = moved_text + "\n\n" + destination_text.lstrip("\n")
            _clear_paragraph(destination)
            destination.add_run(destination_text)
            _remove_paragraph(paragraph)
            paragraphs = _document_table_paragraphs(document)
            break

    code_active = False
    source_markers = ("/* 实验1.1：", "/* 实验1.2：", "/* 实验1.3：")
    section_markers = ("2. 动态分配顺序表及算法", "3. 顺序表的普通合并与有序归并", "运行情况：")
    code_hint = re.compile(
        r"(?:#(?:include|define)|\b(?:typedef|struct|int|void|char|float|double|"
        r"Status|LinkList|ElemType|printf|scanf|malloc|free|return|while|for|if)\b|"
        r"->|[{};])"
    )
    for paragraph in paragraphs:
        text = paragraph.text
        stripped = text.strip()
        if any(marker in text for marker in source_markers):
            code_active = True
        elif any(text.strip().startswith(marker) for marker in section_markers):
            code_active = False
        elif not code_active and code_hint.search(text):
            code_active = True
        if not code_active or not text:
            continue
        if re.fullmatch(r"实验\s*\d+", stripped):
            continue
        for run in paragraph.runs:
            _set_code_run_font(run)


def _restore_wide_table(table, values, allow_header_rebuild=False):
    """Repair a PDF2DOCX table that represents a page of the wide report table."""
    if not table.rows or len(table.columns) != 15:
        return

    header_text = " ".join(cell.text for cell in table.rows[0].cells)
    _set_table_widths(table)
    rebuilt_header = _rebuild_missing_header(table, values) if allow_header_rebuild else False
    if allow_header_rebuild and not rebuilt_header and "课程名称" not in header_text:
        header_row = _insert_row_before(table, 0)
        header_cell = _merge_row(table, 0)
        _write_header_cell(header_cell, values)
        _set_header_row_height(table)
        rebuilt_header = True
        header_text = "课程名称"
    if rebuilt_header:
        header_text = "课程名称"

    if "课程名称" in header_text:
        if not rebuilt_header:
            header_cell = _merge_row(table, 0)
            _write_header_cell(header_cell, values)
            _set_header_row_height(table)

        body_row = table.rows[1]
        body_row.height = Cm(BODY_MIN_HEIGHT_CM)
        body_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if rebuilt_header:
            _set_cell_margins(
                body_row.cells[0],
                start=round(0.20 * 567),
                end=round(0.20 * 567),
                top=round(0.34 * 567),
                bottom=round(0.20 * 567),
            )
        else:
            body_cell = _merge_row(table, 1)
            needs_rewrite = _row_has_multiple_xml_cells(body_row)
            if needs_rewrite:
                body_text = next(
                    (cell.text for cell in body_row.cells if cell.text.strip()),
                    "",
                )
                _write_body_cell(body_cell, body_text)
            else:
                _set_cell_margins(
                    body_cell,
                    start=round(0.20 * 567),
                    end=round(0.20 * 567),
                    top=round(0.34 * 567),
                    bottom=round(0.20 * 567),
                )

    elif len(table.rows) >= 1:
        row_texts = [cell.text.strip() for cell in table.rows[0].cells]
        nonempty = [text for text in row_texts if text]
        if len(nonempty) > 1 and len(set(nonempty)) == 1:
            body_row = table.rows[0]
            needs_rewrite = _row_has_multiple_xml_cells(body_row)
            body_cell = _merge_row(table, 0)
            if needs_rewrite:
                _write_body_cell(body_cell, nonempty[0])
            else:
                _set_cell_margins(
                    body_cell,
                    start=round(0.20 * 567),
                    end=round(0.20 * 567),
                    top=round(0.34 * 567),
                    bottom=round(0.20 * 567),
                )

    evaluation_row = next(
        (row for row in table.rows if any(cell.text.strip() == "教师评价" for cell in row.cells)),
        None,
    )
    if evaluation_row is not None:
        labels = ("教师评价", "优", "", "良", "", "中", "", "及\n格", "", "不及\n格", "", "教师\n签名", "", "日期", "")
        evaluation_row.height = Cm(1.371)
        evaluation_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for index, (cell, label) in enumerate(zip(evaluation_row.cells, labels)):
            alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 9 else WD_ALIGN_PARAGRAPH.CENTER
            # Center the two-character “不及” text block while keeping both
            # lines left-aligned, so “格” starts directly under “不”.
            text_width_cm = 2 * 10.5 * 2.54 / 72
            left_indent_cm = (
                max(0, (REPORT_COLUMNS_CM[index] - text_width_cm) / 2)
                if index == 9
                else 0
            )
            _set_cell_text(
                cell,
                label,
                size=10.5,
                alignment=alignment,
                left_indent_cm=left_indent_cm,
            )
        _set_table_borders(table)


def _restore_header(document, source):
    values = _read_typst_fields(source)
    outer = document.tables[0]
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    if len(outer.columns) == 1 and len(outer.rows) >= 1:
        first_text = " ".join(cell.text for cell in outer.rows[0].cells)
        if any(heading in first_text for heading in BODY_HEADINGS):
            # Some pdf2docx layouts omit the header and return the report body
            # as the first one-column table.  Preserve that body and insert a
            # new header row above it instead of overwriting its only cell.
            if _rebuild_missing_header(outer, values):
                _set_table_borders(outer)
        else:
            _write_header_cell(outer.rows[0].cells[0], values)
            _set_header_row_height(outer)

    for index, table in enumerate(document.tables):
        _restore_wide_table(table, values, allow_header_rebuild=(index == 0))


def _restore_bold_text(document):
    title = "江南大学人工智能与计算机学院实验报告"
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == title:
            for run in paragraph.runs:
                _set_run_bold(run)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for heading in BODY_HEADINGS:
                        for run in paragraph.runs:
                            if heading in run.text:
                                _set_run_bold(run)


def _allow_converted_pages_to_flow(document):
    """Avoid blank pages when restored code slightly changes page height."""
    for section in document.sections[1:]:
        section.start_type = WD_SECTION_START.CONTINUOUS


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_pdf")
    parser.add_argument("output_docx")
    parser.add_argument("source_typ")
    args = parser.parse_args()

    parse(args.input_pdf, args.output_docx)
    document = Document(args.output_docx)
    _collapse_pdf_page_tables(document)
    _restore_header(document, Path(args.source_typ))
    code_texts = _restore_program_blocks(document, Path(args.source_typ))
    _restore_body_font_size(document, code_texts)
    _normalize_code_paragraphs(document)
    _restore_bold_text(document)
    _allow_converted_pages_to_flow(document)
    document.save(args.output_docx)


if __name__ == "__main__":
    main()
